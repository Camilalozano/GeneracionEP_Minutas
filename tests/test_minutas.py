import ast
import io
from pathlib import Path
import subprocess
import tempfile
import unittest
import zipfile
from lxml import etree

import pandas as pd
from docx import Document

import minutas

ROOT = Path(__file__).resolve().parents[1]


def ep_functions(source):
    tree = ast.parse(source)
    wanted = {'replace_text_in_paragraph', 'diligenciar_documento'}
    module = ast.Module(body=[n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in wanted], type_ignores=[])
    scope = {'Document': Document}
    exec(compile(module, '<ep>', 'exec'), scope)
    return scope['diligenciar_documento']


def fixture():
    data = {key: 'Dato de prueba' for key in (*minutas.CONTRATO, *minutas.INSTITUCION)}
    data.update(minutas.DEFAULTS)
    data.update(contratista='PERSONA DE PRUEBA', tipo_contrato='Prestación de servicios profesionales',
                codigo_objeto='COD-PRUEBA', objeto_contrato='Prestar servicios de análisis de información.',
                numero_contrato='001', anio_contrato='2026', cdp_fecha='15/01/2026',
                contratista_documento='0012345', cdp_numero='0009')
    return data


class Progress:
    def progress(self, value):
        assert 0 <= value <= 1
    def text(self, value):
        pass


class MinutasTest(unittest.TestCase):
    def test_all_parts_resolved_and_source_intact(self):
        before = minutas.TEMPLATE.read_bytes()
        output = minutas.generar_minuta(fixture())
        with zipfile.ZipFile(io.BytesIO(output)) as z:
            xml = ''.join(''.join(etree.fromstring(z.read(n)).itertext()) for n in z.namelist() if n.startswith('word/') and n.endswith('.xml'))
        self.assertNotIn('{{', xml)
        self.assertIn('ATENEA-001-2026', xml)
        self.assertIn('0012345', xml)
        self.assertIn('15 de enero de 2026', xml)
        self.assertIn('09/12/2025', xml)
        self.assertEqual(before, minutas.TEMPLATE.read_bytes())

    def test_split_runs_and_table_header(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run('Antes ').italic = True
        p.add_run('{{contr').bold = True
        p.add_run('atista}} después')
        doc.sections[0].header.add_paragraph('{{numero_contrato}}')
        doc.add_table(rows=1, cols=1).cell(0,0).text = '{{cdp_numero}}'
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        result = Document(io.BytesIO(minutas.generar_minuta(fixture(), stream)))
        self.assertEqual(result.paragraphs[0].text, 'Antes PERSONA DE PRUEBA después')
        self.assertTrue(result.paragraphs[0].runs[0].italic)
        self.assertEqual(result.tables[0].cell(0,0).text, '0009')
        self.assertIn('001', '\n'.join(p.text for p in result.sections[0].header.paragraphs))

    def test_reject_missing_wrong_date_and_unsupported_contract(self):
        for key, value in [('contratista', ''), ('cdp_fecha', '31/02/2026'),
                           ('tipo_contrato', 'Compraventa'), ('anio_contrato', '2026.0'),
                           ('numero_contrato', 'ATENEA-001-2026'), ('ordenador_nombre', float('nan'))]:
            with self.subTest(key=key):
                data = fixture(); data[key] = value
                with self.assertRaises(ValueError):
                    minutas.generar_minuta(data)

    def test_ep_engine_and_form_sections_unchanged(self):
        original = subprocess.check_output(['git', 'show', '7ebb910e1232d94e72183f560f3fd479e6d05f45:streamlit_app.py'], cwd=ROOT).decode('utf-8')
        current = (ROOT/'streamlit_app.py').read_text(encoding='utf-8')
        names = ['diligenciar_documento', 'replace_text_in_paragraph', 'validar_formulario',
                 'normalizar_registro_estudios', 'guardar_estudios_previos_supabase',
                 'construir_dataframe_desde_formulario', 'generar_documentos']
        for name in names:
            nodes = [next(n for n in ast.parse(src).body if isinstance(n, ast.FunctionDef) and n.name == name)
                     for src in (original, current)]
            self.assertEqual(ast.dump(nodes[0]), ast.dump(nodes[1]), name)
        start = '        st.markdown("#### 1) Identificación del proceso")'
        end = '        cargar_registro = st.form_submit_button'
        self.assertEqual(original.split(start)[1].split(end)[0].rstrip(),
                         current.split(start)[1].split('        try:\n            configuracion_minuta')[0].rstrip())
        self.assertEqual(subprocess.check_output(['git', 'show', '7ebb910e1232d94e72183f560f3fd479e6d05f45:Plantilla.docx'], cwd=ROOT),
                         (ROOT/'Plantilla.docx').read_bytes())

    def test_pairs_and_storage_failure_preserve_both_outputs(self):
        df = pd.DataFrame([fixture(), fixture()], index=[10,20])
        ep_fn = ep_functions((ROOT/'streamlit_app.py').read_text(encoding='utf-8'))
        def fail(*args):
            raise RuntimeError('unavailable')
        result, count, errors, saved = minutas.generar_pares(
            df, io.BytesIO((ROOT/'Plantilla.docx').read_bytes()), Progress(), Progress(), 'Prueba', ep_fn, fail, fail)
        self.assertEqual((count, saved, len(errors)), (4, 0, 4))
        with zipfile.ZipFile(result) as z:
            self.assertEqual(len(z.namelist()), 4)
            ep = z.read(z.namelist()[0])
        expected = io.BytesIO()
        ep_fn(io.BytesIO((ROOT/'Plantilla.docx').read_bytes()), df.iloc[0]).save(expected)
        with zipfile.ZipFile(io.BytesIO(ep)) as actual, zipfile.ZipFile(expected) as reference:
            self.assertEqual(actual.read('word/document.xml'), reference.read('word/document.xml'))

    def test_invalid_pair_no_upload_or_partial_zip(self):
        data = fixture(); data['cdp_numero'] = ''
        def unexpected(*args):
            self.fail('No generation or persistence should happen before validation')
        result, count, errors, saved = minutas.generar_pares(pd.DataFrame([data]), None,
            Progress(), Progress(), '', unexpected, unexpected, unexpected)
        self.assertEqual(count, 0)
        self.assertTrue(errors)

    def test_private_registry_and_schema_preflight(self):
        calls = []
        class Response:
            def raise_for_status(self):
                pass
        def request(*args, **kwargs):
            calls.append((args, kwargs))
            return Response()
        record = minutas.guardar_minuta_supabase(fixture(), 'pair-test', b'minuta', b'ep',
            {'storage_path':'estudios/ep.docx'}, 'Test', url='https://example.test',
            key='sb_secret_test', bucket='private', campos_ep=['contratista','objeto_contrato'], request=request)
        self.assertEqual(len(calls),3)
        self.assertEqual(calls[0][0][0], 'GET')
        self.assertEqual(record['estudio_previo_storage_path'], 'estudios/ep.docx')
        self.assertEqual(record['estudio_previo_snapshot']['contratista'], fixture()['contratista'])
        self.assertNotIn('sb_secret_test', str(record))


if __name__ == '__main__':
    unittest.main()
